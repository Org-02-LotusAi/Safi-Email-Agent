"""Attachment handling (doc §24). Mandatory types: TXT, PDF, DOCX. Unsupported types
must never crash the agent; parsing failures are captured, not raised.
"""

from __future__ import annotations

import io

from .models import ParsedAttachment

SUPPORTED_EXTS = {"txt", "pdf", "docx", "eml"}
_MANDATORY = {"txt", "pdf", "docx"}  # doc §6/§24


def _ext(filename: str) -> str:
    name = (filename or "").lower()
    _, _, e = name.rpartition(".")
    return e if _ else ""


def _extract_txt(data: bytes) -> str:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader  # imported here so a missing/broken pypdf can't block import

    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text:
            parts.append(text)
    return "\n".join(parts).strip()


def _extract_docx(data: bytes) -> str:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text for cell in row.cells if cell.text]
            if cells:
                parts.append("\t".join(cells))
    return "\n".join(parts).strip()


def _extract_eml(data: bytes) -> str:
    # Forwarded .eml is acceptable dynamic context (doc §22). Extract its plain text.
    from email import message_from_bytes
    from email.policy import default

    msg = message_from_bytes(data, policy=default)
    body = msg.get_body(preferencelist=("plain", "html"))
    if body is None:
        return ""
    text = body.get_content()
    subject = msg.get("Subject", "")
    sender = msg.get("From", "")
    header = f"[Forwarded message] From: {sender} Subject: {subject}\n"
    return (header + text).strip()


def extract(filename: str, data: bytes) -> ParsedAttachment:
    """Extract useful text from an attachment. Never raises for content problems.

    Returns a :class:`ParsedAttachment`; ``supported`` is False for unknown types
    (their content is ignored but the agent keeps running — doc §24, §37/TEST 21).
    """
    ext = _ext(filename)
    ctype = {
        "txt": "text/plain",
        "pdf": "application/pdf",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "eml": "message/rfc822",
    }.get(ext, "application/octet-stream")

    if ext not in SUPPORTED_EXTS:
        return ParsedAttachment(
            filename=filename, content_type=ctype, data=data,
            supported=False, extracted_text=None,
            error=f"unsupported attachment type: .{ext or '(none)'}",
        )

    extractor = {
        "txt": _extract_txt,
        "pdf": _extract_pdf,
        "docx": _extract_docx,
        "eml": _extract_eml,
    }[ext]

    try:
        text = extractor(data)
        return ParsedAttachment(
            filename=filename, content_type=ctype, data=data,
            supported=True, extracted_text=text, error=None,
        )
    except Exception as exc:  # attachment parsing failure (doc §37) — capture, don't crash
        return ParsedAttachment(
            filename=filename, content_type=ctype, data=data,
            supported=True, extracted_text=None,
            error=f"failed to extract .{ext}: {exc}",
        )
